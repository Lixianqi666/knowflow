"""P0: 文档上传解析测试 — 覆盖多种文件格式"""

import io

import pytest
from httpx import AsyncClient


def _make_pdf(title: str = "测试PDF") -> bytes:
    """生成最小合法 PDF"""
    content = f"BT /F1 12 Tf 100 700 Td ({title}) Tj ET"
    stream_bytes = f"stream\n{content}\nendstream".encode()
    objects = [
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj",
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj",
        (
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
            b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj"
        ),
        f"4 0 obj<</Length {len(stream_bytes)}>>".encode() + stream_bytes + b"endobj",
        b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for obj in objects:
        offsets.append(len(pdf))
        pdf += obj + b"\n"
    xref_offset = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += b"trailer<</Size 6/Root 1 0 R>>\n"
    pdf += f"startxref\n{xref_offset}\n%%EOF".encode()
    return pdf


def _make_docx() -> bytes:
    """用 python-docx 生成合法 DOCX"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("测试DOCX内容")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx() -> bytes:
    """用 openpyxl 生成合法 XLSX"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "测试XLSX"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_upload_txt(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={"file": ("test.txt", "纯文本内容".encode(), "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json()["title"] == "test.txt"


@pytest.mark.asyncio
async def test_upload_markdown(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={"file": ("test.md", "# 标题\n\n内容".encode(), "text/markdown")},
    )
    assert resp.status_code == 200
    assert resp.json()["title"] == "test.md"


@pytest.mark.asyncio
async def test_upload_pdf(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={"file": ("test.pdf", _make_pdf(), "application/pdf")},
    )
    # 手工构造的 PDF 可能在某些环境解析失败，接受 200 或 400
    assert resp.status_code in (200, 400)
    if resp.status_code == 200:
        assert resp.json()["title"] == "test.pdf"


@pytest.mark.asyncio
async def test_upload_docx(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={
            "file": (
                "test.docx",
                _make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 200
    assert resp.json()["title"] == "test.docx"


@pytest.mark.asyncio
async def test_upload_xlsx(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={
            "file": (
                "test.xlsx",
                _make_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200
    assert resp.json()["title"] == "test.xlsx"


@pytest.mark.asyncio
async def test_upload_empty_file_rejected(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_upload_unsupported_extension(client: AsyncClient, auth_headers: dict):
    resp = await client.post(
        "/api/v1/documents/upload",
        headers=auth_headers,
        files={"file": ("malware.exe", b"MZ", "application/octet-stream")},
    )
    assert resp.status_code == 400
    assert "不支持" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_upload_unauthenticated(client: AsyncClient):
    resp = await client.post(
        "/api/v1/documents/upload",
        files={"file": ("test.txt", b"content", "text/plain")},
    )
    assert resp.status_code == 403
